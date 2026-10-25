import types

import pytest

import rag
import server


class FakeClient:
    def __init__(self, fail=False):
        self.fail = fail
        self.auth_tokens = types.SimpleNamespace(create=self._create)

    def _create(self, config):
        if self.fail:
            raise RuntimeError("boom")
        return types.SimpleNamespace(name="tok123")


async def test_token_ok(aiohttp_client):
    c = await aiohttp_client(server.make_app(FakeClient()))
    r = await c.post("/api/token")
    assert r.status == 200
    assert (await r.json())["token"] == "tok123"


async def test_token_error(aiohttp_client):
    c = await aiohttp_client(server.make_app(FakeClient(fail=True)))
    r = await c.post("/api/token")
    assert r.status == 500
    assert "error" in await r.json()


async def test_rate_limit(aiohttp_client):
    c = await aiohttp_client(server.make_app(FakeClient()))
    for _ in range(server.RATE_LIMIT):
        assert (await c.post("/api/token")).status == 200
    assert (await c.post("/api/token")).status == 429


async def test_rag_endpoint_finds_and_misses(aiohttp_client):
    c = await aiohttp_client(server.make_app(FakeClient()))
    hit = await (await c.post("/api/rag", json={"query": "how to prompt Cursor"})).json()
    assert any("Cursor" in r["text"] for r in hit["results"])
    miss = await (await c.post("/api/rag", json={"query": "zzzxqv plumbing"})).json()
    assert miss["results"] == []


async def test_rag_bad_input(aiohttp_client):
    c = await aiohttp_client(server.make_app(FakeClient()))
    assert (await c.post("/api/rag", json={})).status == 400
    assert (await c.post("/api/rag", data="not json")).status == 400


def test_index_ranks_relevant_chunk_first():
    idx = rag.Index(rag.load_chunks())
    top = idx.search("PostgreSQL relational database")[0]
    assert "PostgreSQL" in top["text"]


STEPS = [{"title": f"Step {i + 1} - Demo", "prompt": "Role: You are a backend dev.\nGoal: build it\n- first point\n1. numbered line"} for i in range(6)]


async def test_doc_ok(aiohttp_client):
    import io

    from docx import Document

    c = await aiohttp_client(server.make_app(FakeClient()))
    r = await c.post("/api/doc", json={"ai_tool": "Cursor", "steps": STEPS})
    assert r.status == 200
    assert "wordprocessingml" in r.headers["Content-Type"]
    assert "NoWork-prompts.docx" in r.headers["Content-Disposition"]
    doc = Document(io.BytesIO(await r.read()))
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings == [s["title"] for s in STEPS]
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Written for Cursor" in text and "first point" in text and "1. numbered line" in text


@pytest.mark.parametrize(
    "body",
    [
        {},
        {"steps": []},
        {"steps": "nope"},
        {"steps": [{"title": "x", "prompt": "  "}]},
        {"steps": [{"title": "x"}]},
        {"steps": [{"title": "x", "prompt": "y" * 20001}]},
        {"steps": [{"title": "t", "prompt": "p"}] * 13},
    ],
)
async def test_doc_rejects_bad_input(aiohttp_client, body):
    c = await aiohttp_client(server.make_app(FakeClient()))
    assert (await c.post("/api/doc", json=body)).status == 400


async def test_doc_rejects_non_json(aiohttp_client):
    c = await aiohttp_client(server.make_app(FakeClient()))
    assert (await c.post("/api/doc", data="not json")).status == 400


@pytest.mark.parametrize(
    "query, must_contain",
    [
        ("operating system is too large for this beta version", "too large for this beta"),
        ("should I use SQL or NoSQL database", "SQL versus NoSQL"),
        ("MongoDB document database", "MongoDB is a document database"),
        ("Angular framework enterprise", "Angular is a full framework"),
        ("authentication OAuth JWT roles", "Authentication proves who the user is"),
        ("what requirements questions to ask about database", "Ask whether the user prefers SQL"),
        ("Flutter or React Native mobile app", "Flutter builds iOS and Android"),
    ],
)
def test_new_knowledge_is_retrievable(query, must_contain):
    results = rag.Index(rag.load_chunks()).search(query)
    assert any(must_contain in r["text"] for r in results), [r["text"][:50] for r in results]
